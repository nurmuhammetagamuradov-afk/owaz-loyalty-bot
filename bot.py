# -*- coding: utf-8 -*-
"""
Owaz Coffee — бот карты лояльности.
Гость регистрируется (имя + дата рождения), получает карту с QR-кодом.
Кассир сканирует QR или вводит номер карты и начисляет визит.
"""

import os
import io
import re
import sqlite3
import logging
import asyncio
from datetime import datetime

import qrcode

from aiogram import Bot, Dispatcher, F
from aiogram.filters import CommandStart, Command
from aiogram.types import (
    Message,
    CallbackQuery,
    BufferedInputFile,
    InlineKeyboardMarkup,
    InlineKeyboardButton,
    ReplyKeyboardMarkup,
    KeyboardButton,
    ReplyKeyboardRemove,
)

import card_design

logging.basicConfig(level=logging.INFO)

# ---------- НАСТРОЙКИ ----------

BOT_TOKEN = os.environ.get("BOT_TOKEN", "")
BOT_USERNAME = os.environ.get("BOT_USERNAME", "Owaz_Coffee_Bot")
ADMIN_ID = int(os.environ.get("ADMIN_ID", "0"))
CAFE_NAME = os.environ.get("CAFE_NAME", "OWAZ")

DB_DIR = os.environ.get("DB_DIR", "/data")
if not os.path.isdir(DB_DIR):
    DB_DIR = "."
DB_PATH = os.path.join(DB_DIR, "owaz_loyalty.db")

# Уровни: (минимум визитов, название, скидка %)
LEVELS = [
    (100, "VIP", 20),
    (60, "ЗАВСЕГДАТАЙ", 15),
    (30, "ПОСТОЯННЫЙ", 10),
    (10, "ДРУГ", 5),
    (0, "ГОСТЬ", 0),
]

# Шаги регистрации
STEP_DONE, STEP_NAME, STEP_BIRTH = 0, 1, 2


def get_level(visits: int):
    for min_v, name, disc in LEVELS:
        if visits >= min_v:
            return name, disc
    return "ГОСТЬ", 0


def next_level_info(visits: int):
    ups = sorted([l for l in LEVELS if l[0] > visits], key=lambda x: x[0])
    if not ups:
        return None
    min_v, name, disc = ups[0]
    return min_v - visits, name, disc


# ---------- БАЗА ДАННЫХ ----------

def db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    return conn


def init_db():
    conn = db()
    c = conn.cursor()
    c.execute("""
        CREATE TABLE IF NOT EXISTS guests (
            tg_id INTEGER PRIMARY KEY,
            name TEXT,
            phone TEXT,
            visits INTEGER DEFAULT 0,
            created_at TEXT
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS staff (
            tg_id INTEGER PRIMARY KEY,
            name TEXT,
            added_at TEXT
        )
    """)
    c.execute("""
        CREATE TABLE IF NOT EXISTS visits_log (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            guest_id INTEGER,
            staff_id INTEGER,
            change INTEGER,
            created_at TEXT
        )
    """)
    # Новые колонки для уже существующих баз
    have = {r["name"] for r in c.execute("PRAGMA table_info(guests)")}
    if "birthday" not in have:
        c.execute("ALTER TABLE guests ADD COLUMN birthday TEXT")
    if "reg_step" not in have:
        c.execute("ALTER TABLE guests ADD COLUMN reg_step INTEGER DEFAULT 0")
    conn.commit()
    conn.close()


def get_guest(tg_id):
    conn = db()
    row = conn.execute("SELECT * FROM guests WHERE tg_id=?", (tg_id,)).fetchone()
    conn.close()
    return row


def create_guest(tg_id):
    conn = db()
    conn.execute(
        "INSERT OR IGNORE INTO guests (tg_id, name, visits, created_at, reg_step) "
        "VALUES (?,?,0,?,?)",
        (tg_id, "", datetime.now().isoformat(), STEP_NAME),
    )
    conn.commit()
    conn.close()


def update_guest(tg_id, **fields):
    if not fields:
        return
    sets = ", ".join(f"{k}=?" for k in fields)
    conn = db()
    conn.execute(f"UPDATE guests SET {sets} WHERE tg_id=?",
                 (*fields.values(), tg_id))
    conn.commit()
    conn.close()


def change_visits(guest_id, staff_id, delta):
    conn = db()
    conn.execute("UPDATE guests SET visits = MAX(0, visits + ?) WHERE tg_id=?",
                 (delta, guest_id))
    conn.execute(
        "INSERT INTO visits_log (guest_id, staff_id, change, created_at) VALUES (?,?,?,?)",
        (guest_id, staff_id, delta, datetime.now().isoformat()),
    )
    conn.commit()
    row = conn.execute("SELECT visits FROM guests WHERE tg_id=?", (guest_id,)).fetchone()
    conn.close()
    return row["visits"] if row else 0


def is_staff(tg_id):
    if tg_id == ADMIN_ID:
        return True
    conn = db()
    row = conn.execute("SELECT 1 FROM staff WHERE tg_id=?", (tg_id,)).fetchone()
    conn.close()
    return row is not None


def add_staff(tg_id, name):
    conn = db()
    conn.execute("INSERT OR REPLACE INTO staff (tg_id, name, added_at) VALUES (?,?,?)",
                 (tg_id, name, datetime.now().isoformat()))
    conn.commit()
    conn.close()


def remove_staff(tg_id):
    conn = db()
    conn.execute("DELETE FROM staff WHERE tg_id=?", (tg_id,))
    conn.commit()
    conn.close()


def stats():
    conn = db()
    total = conn.execute("SELECT COUNT(*) c FROM guests WHERE reg_step=0").fetchone()["c"]
    visits = conn.execute("SELECT COALESCE(SUM(visits),0) s FROM guests").fetchone()["s"]
    today = datetime.now().strftime("%Y-%m-%d")
    today_v = conn.execute(
        "SELECT COUNT(*) c FROM visits_log WHERE change>0 AND created_at LIKE ?",
        (today + "%",)).fetchone()["c"]
    conn.close()
    return total, visits, today_v


# ---------- КАРТА ----------

def render_card(tg_id, name, visits):
    level, discount = get_level(visits)
    nxt = next_level_info(visits)
    next_info = None
    if nxt:
        need, nname, ndisc = nxt
        prev = max([l[0] for l in LEVELS if l[0] <= visits] or [0])
        total = (visits + need) - prev
        next_info = (need, nname, ndisc, visits - prev, total)

    qr = qrcode.QRCode(box_size=10, border=1)
    qr.add_data(f"https://t.me/{BOT_USERNAME}?start=g{tg_id}")
    qr.make(fit=True)
    qr_img = qr.make_image(fill_color="black", back_color="white")

    return card_design.render(tg_id, name, visits, level, discount,
                              next_info, CAFE_NAME, qr_img)


# ---------- КЛАВИАТУРЫ ----------

def guest_kb():
    return ReplyKeyboardMarkup(
        keyboard=[
            [KeyboardButton(text="💳 Meniň kartym / Моя карта")],
            [KeyboardButton(text="ℹ️ Şertler / Условия")],
        ],
        resize_keyboard=True,
    )


def staff_panel_kb(guest_id):
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="➕ Начислить визит", callback_data=f"add:{guest_id}")],
        [InlineKeyboardButton(text="➖ Отменить визит", callback_data=f"sub:{guest_id}")],
    ])


# ---------- БОТ ----------

bot = Bot(token=BOT_TOKEN)
dp = Dispatcher()


async def show_guest_panel(message: Message, guest_id: int):
    g = get_guest(guest_id)
    if not g or g["reg_step"] != STEP_DONE:
        await message.answer("Гость не найден. Проверьте номер карты.")
        return
    level, disc = get_level(g["visits"])
    await message.answer(
        f"👤 <b>{g['name']}</b>\n"
        f"Визиты: <b>{g['visits']}</b>\n"
        f"Уровень: <b>{level}</b>\n"
        f"Скидка: <b>{disc}%</b>\n"
        f"Карта: <code>{guest_id}</code>",
        parse_mode="HTML",
        reply_markup=staff_panel_kb(guest_id),
    )


async def send_card(message: Message, tg_id: int):
    g = get_guest(tg_id)
    png = render_card(tg_id, g["name"] or "Гость", g["visits"])
    await message.answer_photo(
        BufferedInputFile(png, filename="card.png"),
        caption="👆 Покажите QR-код кассиру / QR kody kassira görkeziň",
        reply_markup=guest_kb(),
    )


@dp.message(CommandStart())
async def cmd_start(message: Message):
    args = message.text.split(maxsplit=1)
    payload = args[1].strip() if len(args) > 1 else ""

    # Кассир отсканировал QR гостя
    if payload.startswith("g") and payload[1:].isdigit():
        guest_id = int(payload[1:])
        if is_staff(message.from_user.id) and guest_id != message.from_user.id:
            await show_guest_panel(message, guest_id)
            return

    g = get_guest(message.from_user.id)
    if g and g["reg_step"] == STEP_DONE:
        await send_card(message, message.from_user.id)
        return

    create_guest(message.from_user.id)
    update_guest(message.from_user.id, reg_step=STEP_NAME)
    await message.answer(
        f"Salam! 👋 Добро пожаловать в клуб гостей {CAFE_NAME}!\n\n"
        "Чтобы оформить карту, ответьте на два вопроса.\n\n"
        "<b>Как вас зовут?</b>\n"
        "Напишите имя и фамилию.",
        parse_mode="HTML",
        reply_markup=ReplyKeyboardRemove(),
    )


def parse_birthday(text):
    """Принимает 01.02.1995, 01/02/1995, 1.2.95 и т.п."""
    m = re.match(r"^\s*(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{2,4})\s*$", text)
    if not m:
        return None
    d, mo, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
    if y < 100:
        y += 1900 if y > 30 else 2000
    try:
        dt = datetime(y, mo, d)
    except ValueError:
        return None
    if dt.year < 1920 or dt > datetime.now():
        return None
    return dt.strftime("%Y-%m-%d")


@dp.message(F.text, ~F.text.startswith("/"))
async def on_text(message: Message):
    uid = message.from_user.id
    text = (message.text or "").strip()
    g = get_guest(uid)

    # --- Регистрация ---
    if g and g["reg_step"] == STEP_NAME:
        name = text[:40]
        if len(name) < 2:
            await message.answer("Слишком короткое имя. Напишите имя и фамилию.")
            return
        update_guest(uid, name=name, reg_step=STEP_BIRTH)
        await message.answer(
            f"Приятно познакомиться, {name}!\n\n"
            "<b>Когда у вас день рождения?</b>\n"
            "Напишите дату в формате ДД.ММ.ГГГГ — например 14.03.1995.\n"
            "Мы дарим подарки именинникам 🎁",
            parse_mode="HTML",
        )
        return

    if g and g["reg_step"] == STEP_BIRTH:
        bd = parse_birthday(text)
        if not bd:
            await message.answer(
                "Не понял дату. Напишите в формате ДД.ММ.ГГГГ — например 14.03.1995."
            )
            return
        update_guest(uid, birthday=bd, reg_step=STEP_DONE)
        await message.answer(
            "Готово! Ваша карта оформлена 🎉\n\n"
            "Показывайте QR-код кассиру при каждом визите — "
            "визиты копятся, скидка растёт.",
            reply_markup=guest_kb(),
        )
        await send_card(message, uid)
        return

    # --- Кнопки гостя ---
    if "Моя карта" in text or "kartym" in text:
        if not g or g["reg_step"] != STEP_DONE:
            await message.answer("Сначала завершите регистрацию: отправьте /start")
            return
        await send_card(message, uid)
        return

    if "Условия" in text or "Şertler" in text:
        t = f"<b>{CAFE_NAME} — карта лояльности</b>\n\n"
        for min_v, name, disc in sorted(LEVELS, key=lambda x: x[0]):
            if disc == 0:
                t += f"• <b>{name}</b> — от {min_v} визитов — копим\n"
            else:
                t += f"• <b>{name}</b> — от {min_v} визитов — скидка <b>{disc}%</b>\n"
        t += ("\nОдин визит = 1 отметка.\nСкидка применяется автоматически.\n"
              "В день рождения — подарок от кофейни 🎁")
        await message.answer(t, parse_mode="HTML")
        return

    # --- Кассир ввёл номер карты ---
    if is_staff(uid) and text.isdigit() and len(text) >= 5:
        await show_guest_panel(message, int(text))
        return

    if is_staff(uid):
        await message.answer(
            "Отсканируйте QR гостя или отправьте номер карты (цифры под QR)."
        )


@dp.callback_query(F.data.startswith("add:"))
async def cb_add(call: CallbackQuery):
    if not is_staff(call.from_user.id):
        await call.answer("Нет доступа", show_alert=True)
        return
    guest_id = int(call.data.split(":")[1])
    old = get_guest(guest_id)
    if not old:
        await call.answer("Гость не найден", show_alert=True)
        return
    old_level, _ = get_level(old["visits"])
    new_visits = change_visits(guest_id, call.from_user.id, +1)
    level, disc = get_level(new_visits)

    await call.message.edit_text(
        f"✅ Визит начислен\n\n"
        f"👤 <b>{old['name']}</b>\n"
        f"Визиты: <b>{new_visits}</b>\n"
        f"Уровень: <b>{level}</b>\n"
        f"Скидка: <b>{disc}%</b>",
        parse_mode="HTML",
        reply_markup=staff_panel_kb(guest_id),
    )
    await call.answer("Визит +1")

    try:
        if level != old_level:
            await bot.send_message(
                guest_id,
                f"🎉 Поздравляем! Новый уровень: <b>{level}</b>\n"
                f"Ваша скидка теперь <b>{disc}%</b>",
                parse_mode="HTML")
        else:
            await bot.send_message(
                guest_id,
                f"✅ Визит засчитан.\nВсего: <b>{new_visits}</b> · Скидка: <b>{disc}%</b>",
                parse_mode="HTML")
    except Exception as e:
        logging.warning(f"Не смог уведомить гостя: {e}")


@dp.callback_query(F.data.startswith("sub:"))
async def cb_sub(call: CallbackQuery):
    if not is_staff(call.from_user.id):
        await call.answer("Нет доступа", show_alert=True)
        return
    guest_id = int(call.data.split(":")[1])
    g = get_guest(guest_id)
    new_visits = change_visits(guest_id, call.from_user.id, -1)
    level, disc = get_level(new_visits)
    await call.message.edit_text(
        f"↩️ Визит отменён\n\n"
        f"👤 <b>{g['name']}</b>\n"
        f"Визиты: <b>{new_visits}</b>\n"
        f"Уровень: <b>{level}</b> · {disc}%",
        parse_mode="HTML",
        reply_markup=staff_panel_kb(guest_id),
    )
    await call.answer("Визит -1")


# ---------- ОТЧЁТ В EXCEL ----------

def build_report():
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter

    conn = db()
    guests = conn.execute("""
        SELECT g.tg_id, g.name, g.visits, g.birthday, g.created_at,
               (SELECT MAX(created_at) FROM visits_log
                 WHERE guest_id = g.tg_id AND change > 0) AS last_visit
        FROM guests g WHERE g.reg_step = 0
        ORDER BY g.visits DESC, g.name
    """).fetchall()
    logs = conn.execute("""
        SELECT v.created_at, v.guest_id, g.name, v.change
        FROM visits_log v LEFT JOIN guests g ON g.tg_id = v.guest_id
        ORDER BY v.created_at DESC LIMIT 5000
    """).fetchall()
    conn.close()

    head_font = Font(name="Arial", bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="8A6D3B")
    body_font = Font(name="Arial")

    def ru_date(iso):
        if not iso:
            return "—"
        try:
            return datetime.strptime(iso[:10], "%Y-%m-%d").strftime("%d.%m.%Y")
        except ValueError:
            return iso[:10]

    wb = Workbook()

    ws = wb.active
    ws.title = "Гости"
    ws.append(["Имя гостя", "День рождения", "Визиты", "Уровень", "Скидка %",
               "Регистрация", "Последний визит", "Номер карты"])
    for g in guests:
        level, disc = get_level(g["visits"])
        ws.append([g["name"], ru_date(g["birthday"]), g["visits"], level, disc,
                   ru_date(g["created_at"]),
                   (g["last_visit"] or "—")[:16].replace("T", " "),
                   g["tg_id"]])

    ws2 = wb.create_sheet("Журнал визитов")
    ws2.append(["Дата и время", "Гость", "Операция", "Номер карты"])
    for r in logs:
        ws2.append([(r["created_at"] or "")[:16].replace("T", " "),
                    r["name"] or "—",
                    "Визит +1" if r["change"] > 0 else "Отмена -1",
                    r["guest_id"]])

    ws3 = wb.create_sheet("Дни рождения")
    ws3.append(["Дата", "Имя гостя", "Визиты", "Уровень"])
    bdays = [g for g in guests if g["birthday"]]
    bdays.sort(key=lambda g: g["birthday"][5:])
    for g in bdays:
        level, _ = get_level(g["visits"])
        ws3.append([ru_date(g["birthday"])[:5], g["name"], g["visits"], level])

    for sheet, widths in ((ws, [26, 14, 9, 16, 10, 14, 18, 14]),
                          (ws2, [20, 26, 14, 14]),
                          (ws3, [10, 26, 9, 16])):
        for i, w in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = w
        for cell in sheet[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="center")
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
        sheet.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


@dp.message(Command("report"))
async def cmd_report(message: Message):
    if message.from_user.id != ADMIN_ID:
        return
    try:
        data = build_report()
    except Exception as e:
        logging.exception("Ошибка отчёта")
        await message.answer(f"Не удалось собрать отчёт: {e}")
        return
    total, visits, today_v = stats()
    stamp = datetime.now().strftime("%Y-%m-%d")
    await message.answer_document(
        BufferedInputFile(data, filename=f"owaz-guests-{stamp}.xlsx"),
        caption=(f"📊 Отчёт на {stamp}\n"
                 f"Гостей: {total} · Визитов всего: {visits} · Сегодня: {today_v}"),
    )


# ---------- АДМИН ----------

@dp.message(Command("reset"))
async def cmd_reset(message: Message):
    """Пройти регистрацию заново (имя и дата рождения). Визиты сохраняются."""
    uid = message.from_user.id
    g = get_guest(uid)
    if not g:
        await message.answer("Вы ещё не зарегистрированы. Отправьте /start")
        return
    update_guest(uid, reg_step=STEP_NAME)
    await message.answer(
        "Давайте заполним карту заново. Визиты сохранятся.\n\n"
        "<b>Как вас зовут?</b>\nНапишите имя и фамилию.",
        parse_mode="HTML",
        reply_markup=ReplyKeyboardRemove(),
    )


@dp.message(Command("id"))
async def cmd_id(message: Message):
    await message.answer(f"Ваш ID: <code>{message.from_user.id}</code>", parse_mode="HTML")


@dp.message(Command("admin"))
async def cmd_admin(message: Message):
    if message.from_user.id != ADMIN_ID:
        return
    total, visits, today_v = stats()
    conn = db()
    staff_rows = conn.execute("SELECT * FROM staff").fetchall()
    conn.close()
    slist = "\n".join(f"• {s['name']} — <code>{s['tg_id']}</code>"
                      for s in staff_rows) or "— нет —"
    await message.answer(
        "<b>Панель администратора</b>\n\n"
        f"Гостей: <b>{total}</b>\n"
        f"Визитов всего: <b>{visits}</b>\n"
        f"Визитов сегодня: <b>{today_v}</b>\n\n"
        f"<b>Кассиры:</b>\n{slist}\n\n"
        "Команды:\n"
        "<code>/report</code> — таблица гостей в Excel\n"
        "<code>/addstaff ID Имя</code> — добавить кассира\n"
        "<code>/delstaff ID</code> — убрать кассира\n"
        "<code>/id</code> — узнать свой Telegram ID\n"
        "<code>/reset</code> — заполнить свою карту заново",
        parse_mode="HTML")


@dp.message(Command("addstaff"))
async def cmd_addstaff(message: Message):
    if message.from_user.id != ADMIN_ID:
        return
    parts = message.text.split(maxsplit=2)
    if len(parts) < 2 or not parts[1].isdigit():
        await message.answer("Формат: /addstaff 123456789 Ислам")
        return
    name = parts[2] if len(parts) > 2 else "Кассир"
    add_staff(int(parts[1]), name)
    await message.answer(f"✅ Кассир {name} добавлен.")


@dp.message(Command("delstaff"))
async def cmd_delstaff(message: Message):
    if message.from_user.id != ADMIN_ID:
        return
    parts = message.text.split()
    if len(parts) < 2 or not parts[1].isdigit():
        await message.answer("Формат: /delstaff 123456789")
        return
    remove_staff(int(parts[1]))
    await message.answer("✅ Кассир удалён.")


async def main():
    init_db()
    logging.info("Бот запускается...")
    await dp.start_polling(bot)


if __name__ == "__main__":
    asyncio.run(main())
